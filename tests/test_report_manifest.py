"""Tests for deterministic report packages and replay verification."""

from __future__ import annotations

import os
from pathlib import Path
import subprocess
import sys
from dataclasses import replace

import pandas as pd
import pytest
from openpyxl import load_workbook

from stock_tool.application.reporting import ReportingService, verify_report_replay
from stock_tool.backtest import BacktestEngine
from stock_tool.data.corporate_actions import CorporateAction, CorporateActionType
from stock_tool.domain.models import Market, Symbol
from stock_tool.reports import ReportData, generate_excel_report, generate_html_report
from stock_tool.reports.models import BenchmarkMetadata
from stock_tool.risk import RiskAlert


def _prices() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "date": ["2024-01-01", "2024-01-02", "2024-01-03", "2024-01-04"],
            "symbol": ["2330"] * 4,
            "open": [100.0, 100.0, 105.0, 103.0],
            "high": [101.0, 106.0, 106.0, 113.0],
            "low": [99.0, 99.0, 102.0, 102.0],
            "close": [100.0, 105.0, 103.0, 112.0],
            "volume": [1000, 1100, 1200, 1300],
        }
    )


def _benchmark() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "date": ["2024-01-01", "2024-01-02", "2024-01-03", "2024-01-04"],
            "close": [100.0, 110.0, 99.0, 108.0],
        }
    )


def _report_data() -> ReportData:
    prices = _prices()
    result = BacktestEngine(initial_cash=10_000).run(
        prices,
        pd.DataFrame([{"date": "2024-01-01", "symbol": "2330", "signal": 1, "quantity": 10}]),
        benchmark=_benchmark(),
    )
    return ReportData(
        symbol="2330",
        market="TWSE",
        analysis_date="2024-01-04",
        price_data=prices,
        backtest_result=result,
        parameters={"strategy": "sample", "commission_rate": 0.001, "tax_rate": 0.003},
    )


def _metadata() -> BenchmarkMetadata:
    return BenchmarkMetadata(
        symbol="0050",
        market="TWSE",
        provider="yfinance",
        provider_symbol="0050.TW",
        source_type="online",
        start_date="2024-01-01",
        end_date="2024-01-04",
        interval="1d",
        row_count=4,
        fetched_at="2024-01-05T00:00:00Z",
        warnings=("Authorization: Bearer manifest-secret",),
    )


def test_same_semantic_inputs_produce_same_manifest_hash_despite_frame_order() -> None:
    service = ReportingService(stocktool_version="1.2.0")
    first = service.build_package(
        _report_data(), benchmark_data=_benchmark(), benchmark_metadata=_metadata()
    )
    reordered = _report_data()
    assert reordered.price_data is not None
    reordered_prices = reordered.price_data.loc[
        :, list(reversed(reordered.price_data.columns))
    ].iloc[::-1]
    second = service.build_package(
        ReportData(
            symbol=reordered.symbol,
            market=reordered.market,
            analysis_date=reordered.analysis_date,
            price_data=reordered_prices,
            backtest_result=reordered.backtest_result,
            parameters=reordered.parameters,
        ),
        benchmark_data=_benchmark().iloc[::-1],
        benchmark_metadata=_metadata(),
    )

    assert first.manifest.manifest_hash == second.manifest.manifest_hash
    assert "manifest-secret" not in first.manifest.canonical_json
    assert "[REDACTED]" in first.manifest.canonical_json


def test_report_package_uses_one_explicit_benchmark_basis_when_metadata_conflicts() -> None:
    package = ReportingService(stocktool_version="1.2.2").build_package(
        _report_data(),
        benchmark_data=_benchmark(),
        benchmark_metadata=BenchmarkMetadata(
            symbol="0050",
            market="TWSE",
            provider="fixture",
            provider_symbol="0050.TW",
            source_type="fixture",
            start_date="2024-01-01",
            end_date="2024-01-04",
            interval="1d",
            row_count=4,
            return_basis="total_return",
            price_policy="adjusted_total_return",
        ),
    )

    assert package.benchmark.available is False
    assert package.benchmark.missing_reason == "benchmark_return_basis_mismatch"
    assert (
        package.data.resolved_backtest_metrics()
        .loc[lambda frame: frame["指標"] == "基準總報酬率", "數值"]
        .iloc[0]
        == "資料不足"
    )
    assert package.manifest.payload["benchmark_alignment"]["missing_reason"] == (
        "benchmark_return_basis_mismatch"
    )


def test_manifest_hashes_corporate_action_audit_and_daily_reconciliation() -> None:
    prices = _prices().assign(market="TWSE")
    split = CorporateAction(
        symbol=Symbol("2330", Market.TWSE),
        action_type=CorporateActionType.SPLIT,
        effective_date="2024-01-03",
        available_date="2024-01-02",
        split_ratio=2.0,
        source="fixture",
    )
    result = BacktestEngine(initial_cash=10_000, corporate_actions=(split,)).run(
        prices,
        pd.DataFrame([{"date": "2024-01-01", "symbol": "2330", "signal": 1, "quantity": 10}]),
    )
    source = ReportData(symbol="2330", market="TWSE", price_data=prices, backtest_result=result)
    service = ReportingService(stocktool_version="1.2.2")
    first = service.build_package(source)

    changed_audit = result.corporate_action_audit.copy(deep=True)
    changed_audit.loc[0, "status"] = "unavailable"
    changed_result = replace(result, corporate_action_audit=changed_audit)
    second = service.build_package(replace(source, backtest_result=changed_result))

    assert (
        first.manifest.payload["output_hashes"]["corporate_action_audit"]
        != second.manifest.payload["output_hashes"]["corporate_action_audit"]
    )
    assert first.manifest.manifest_hash != second.manifest.manifest_hash


def test_report_package_replay_detects_changed_price_and_keeps_inputs_immutable() -> None:
    service = ReportingService(stocktool_version="1.2.0")
    source = _report_data()
    assert source.price_data is not None
    source_prices = source.price_data
    before = source_prices.copy(deep=True)
    package = service.build_package(
        source, benchmark_data=_benchmark(), benchmark_metadata=_metadata()
    )

    assert verify_report_replay(package, source, service=service).matches is True
    changed = source_prices.copy(deep=True)
    changed.loc[0, "close"] = 101.0
    mismatch = verify_report_replay(
        package,
        ReportData(
            symbol=source.symbol,
            market=source.market,
            analysis_date=source.analysis_date,
            price_data=changed,
            backtest_result=source.backtest_result,
            parameters=source.parameters,
        ),
        service=service,
    )

    assert mismatch.matches is False
    assert "price_input_hash" in mismatch.mismatches
    pd.testing.assert_frame_equal(source_prices, before)


def test_strategy_cost_and_private_provider_metadata_change_or_redact_manifest() -> None:
    service = ReportingService(stocktool_version="1.2.0")
    source = _report_data()
    baseline = service.build_package(
        source,
        benchmark_data=_benchmark(),
        provider_metadata={
            "provider": "yfinance",
            "warnings": (
                "API key: private-key",
                "download C:\\Users\\steve\\private-cache\\2330.csv",
            ),
            "attempts": ({"reason": "token=hidden-token"},),
        },
    )
    changed = service.build_package(
        ReportData(
            symbol=source.symbol,
            market=source.market,
            analysis_date=source.analysis_date,
            price_data=source.price_data,
            backtest_result=source.backtest_result,
            parameters={
                **dict(source.parameters or {}),
                "strategy_parameters": {"fast_window": 10, "slow_window": 60},
                "commission_rate": 0.002,
            },
        ),
        benchmark_data=_benchmark(),
    )

    assert baseline.manifest.manifest_hash != changed.manifest.manifest_hash
    assert "private-key" not in baseline.manifest.canonical_json
    assert "hidden-token" not in baseline.manifest.canonical_json
    assert "C:\\Users\\steve" not in baseline.manifest.canonical_json
    assert "[PATH_REDACTED]" in baseline.manifest.canonical_json
    assert "[REDACTED]" in baseline.manifest.canonical_json


def test_excel_and_html_share_manifest_benchmark_and_missing_data_state(tmp_path: Path) -> None:
    service = ReportingService(stocktool_version="1.2.0")
    package = service.build_package(
        _report_data(), benchmark_data=_benchmark(), benchmark_metadata=_metadata()
    )

    excel = generate_excel_report(package, tmp_path / "report.xlsx")
    html = generate_html_report(package, tmp_path / "report.html")

    workbook = load_workbook(excel, read_only=True)
    assert "可重現性 Manifest" in workbook.sheetnames
    assert "基準比較" in workbook.sheetnames
    manifest_values = [
        cell.value for row in workbook["可重現性 Manifest"].iter_rows() for cell in row
    ]
    assert package.manifest.manifest_hash in manifest_values
    html_text = html.read_text(encoding="utf-8")
    assert package.manifest.manifest_hash in html_text
    assert "基準比較" in html_text
    workbook.close()


def test_missing_benchmark_is_displayed_as_insufficient_data_not_zero() -> None:
    package = ReportingService(stocktool_version="1.2.0").build_package(_report_data())

    assert package.benchmark.available is False
    assert package.benchmark.total_return is None
    assert "資料不足" in package.data.resolved_benchmark_comparison().to_string()


def test_excel_external_text_is_always_written_as_literal_text(tmp_path: Path) -> None:
    source = _report_data()
    source = replace(
        source,
        technical_summary='=HYPERLINK("https://evil.invalid","click")',
        risk_alerts=[RiskAlert(code="external", severity="warning", message="@evil")],
        provider_metadata={
            "provider": "+cmd|'/C calc'!A0",
            "warnings": ("-2+3",),
        },
    )
    package = ReportingService(stocktool_version="1.2.0").build_package(
        source,
        benchmark_data=_benchmark(),
    )
    output = generate_excel_report(package, tmp_path / "formula-safety.xlsx")

    workbook = load_workbook(output, data_only=False)
    formula_cells = [
        cell
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.lstrip().startswith(("=", "+", "-", "@"))
    ]
    assert formula_cells == []
    assert all(
        cell.data_type != "f"
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows()
        for cell in row
    )
    workbook.close()


def test_manifest_contains_analysis_technical_and_output_hashes() -> None:
    base = _report_data()
    assert base.backtest_result is not None
    source = replace(
        base,
        technical_indicators=pd.DataFrame(
            {"date": ["2024-01-01", "2024-01-02"], "sma_20": [100.0, 101.0]}
        ),
        trade_log=pd.DataFrame({"date": ["2024-01-02"], "message": ["filled"]}),
        equity_curve=base.backtest_result.equity_curve.copy(),
        risk_alerts=[RiskAlert(code="risk", severity="warning", message="check")],
    )
    package = ReportingService(stocktool_version="1.2.0").build_package(
        source,
        benchmark_data=_benchmark(),
    )

    payload = package.manifest.payload
    assert payload["analysis_date"] == "2024-01-04"
    assert payload["technical_input_hash"]
    assert set(payload["output_hashes"]) == {
        "trade_log",
        "equity_curve",
        "risk_alerts",
        "corporate_action_audit",
        "daily_reconciliation",
        "backtest_metrics",
    }
    assert all(value for value in payload["output_hashes"].values())
    assert "filled" not in package.manifest.canonical_json
    assert "check" not in package.manifest.canonical_json


def test_report_content_changes_change_manifest_hash() -> None:
    service = ReportingService(stocktool_version="1.2.0")
    base = _report_data()
    assert base.backtest_result is not None
    source = replace(
        base,
        technical_indicators=pd.DataFrame({"date": ["2024-01-01"], "sma_20": [100.0]}),
        trade_log=pd.DataFrame({"date": ["2024-01-02"], "message": ["filled"]}),
        equity_curve=base.backtest_result.equity_curve.copy(),
        risk_alerts=[RiskAlert(code="risk", severity="warning", message="check")],
    )
    benchmark = _benchmark()
    baseline = service.build_package(source, benchmark_data=benchmark)
    assert source.equity_curve is not None
    assert source.backtest_result is not None

    variants = [
        replace(source, analysis_date="2024-01-05"),
        replace(
            source,
            technical_indicators=pd.DataFrame({"date": ["2024-01-01"], "sma_20": [101.0]}),
        ),
        replace(source, trade_log=pd.DataFrame({"date": ["2024-01-02"], "message": ["changed"]})),
        replace(
            source,
            equity_curve=source.equity_curve.assign(total_equity=[100.0, 106.0, 103.0, 112.0]),
        ),
        replace(
            source,
            risk_alerts=[RiskAlert(code="risk", severity="warning", message="changed")],
        ),
        replace(
            source,
            backtest_result=replace(
                source.backtest_result,
                metrics=replace(source.backtest_result.metrics, total_return=0.25),
            ),
        ),
    ]

    for variant in variants:
        candidate = service.build_package(variant, benchmark_data=benchmark)
        assert candidate.manifest.manifest_hash != baseline.manifest.manifest_hash


def test_semantic_report_frames_keep_hash_when_rows_and_columns_are_reordered() -> None:
    service = ReportingService(stocktool_version="1.2.0")
    source = replace(
        _report_data(),
        technical_indicators=pd.DataFrame(
            {"date": ["2024-01-01", "2024-01-02"], "sma_20": [100.0, 101.0]}
        ),
        trade_log=pd.DataFrame({"date": ["2024-01-02", "2024-01-03"], "message": ["a", "b"]}),
        risk_alerts=[RiskAlert(code="risk", severity="warning", message="check")],
    )
    assert source.technical_indicators is not None
    assert source.trade_log is not None
    reordered = replace(
        source,
        technical_indicators=source.technical_indicators.loc[:, ["sma_20", "date"]].iloc[::-1],
        trade_log=source.trade_log.iloc[::-1],
    )
    first = service.build_package(source, benchmark_data=_benchmark())
    second = service.build_package(reordered, benchmark_data=_benchmark())
    assert first.manifest.manifest_hash == second.manifest.manifest_hash


def test_same_set_hash_is_stable_across_ten_fresh_python_processes() -> None:
    code = (
        "from stock_tool.application.reporting import canonical_object_hash; "
        "print(canonical_object_hash({'values': {'beta', 'alpha', 'gamma'}}))"
    )
    env = os.environ.copy()
    env["PYTHONPATH"] = str(Path.cwd() / "src")
    hashes = [
        subprocess.check_output([sys.executable, "-c", code], env=env, text=True).strip()
        for _ in range(10)
    ]
    assert len(set(hashes)) == 1


def test_package_uses_one_canonical_benchmark_for_metrics_html_and_manifest(tmp_path: Path) -> None:
    service = ReportingService(stocktool_version="1.2.0")
    package = service.build_package(_report_data(), benchmark_data=_benchmark())
    alternate = _benchmark().assign(close=[100.0, 120.0, 90.0, 130.0])
    alternate_package = service.build_package(_report_data(), benchmark_data=alternate)

    assert package.data.benchmark_comparison is package.benchmark
    assert (
        package.data.resolved_backtest_metrics()
        .loc[lambda frame: frame["指標"] == "基準總報酬率", "數值"]
        .iloc[0]
        == "8.00%"
    )
    assert (
        alternate_package.data.resolved_backtest_metrics()
        .loc[lambda frame: frame["指標"] == "基準總報酬率", "數值"]
        .iloc[0]
        == "30.00%"
    )
    html = generate_html_report(package, tmp_path / "canonical-benchmark.html")
    html_text = html.read_text(encoding="utf-8")
    assert "8.00%" in html_text
    assert "30.00%" not in html_text
    assert float(
        package.manifest.payload["output_summary"]["benchmark_total_return"]
    ) == pytest.approx(0.08)


def test_timezone_aware_and_naive_trade_dates_align_without_type_error() -> None:
    equity = _equity_curve_with_timezone()
    benchmark = _benchmark_with_timezone()
    source = replace(_report_data(), equity_curve=equity)
    package = ReportingService(stocktool_version="1.2.0").build_package(
        source,
        benchmark_data=benchmark,
    )
    assert package.benchmark.available is True


def _equity_curve_with_timezone() -> pd.DataFrame:
    return pd.DataFrame(
        {"date": pd.date_range("2024-01-01", periods=4), "total_equity": [100, 105, 103, 112]}
    )


def _benchmark_with_timezone() -> pd.DataFrame:
    return pd.DataFrame(
        {"date": pd.date_range("2024-01-01", periods=4, tz="UTC"), "close": [100, 110, 99, 108]}
    )

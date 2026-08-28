from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from stock_tool.backtest import BacktestEngine
from stock_tool.reports import (
    SURVIVORSHIP_BIAS_WARNING,
    TRAILING_STOP_DAILY_MODEL_NOTICE,
    ReportData,
    generate_excel_report,
    generate_html_report,
)
from stock_tool.reports.excel import SHEET_NAMES
from stock_tool.reports.html import render_html_report
from stock_tool.risk import RiskAlert


def _price_data() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "date": ["2024-01-01", "2024-01-02", "2024-01-03"],
            "symbol": ["2330", "2330", "2330"],
            "open": [100.0, 101.0, 102.0],
            "high": [101.0, 102.0, 103.0],
            "low": [99.0, 100.0, 101.0],
            "close": [100.0, 102.0, 101.0],
            "volume": [1000, 1200, 1100],
        }
    )


def _report_data() -> ReportData:
    prices = _price_data()
    indicators = prices.assign(sma_20=[None, None, 101.0], rsi_14=[None, None, 55.0])
    fundamentals = pd.DataFrame(
        {
            "symbol": ["2330"],
            "total_score": [75.0],
            "growth_score": [20.0],
            "profitability_score": [20.0],
            "financial_safety_score": [15.0],
            "valuation_score": [12.0],
            "cashflow_score": [8.0],
            "strengths": ["營收年增率尚可接受"],
            "weaknesses": ["股價淨值比偏高"],
            "missing_data": [""],
            "risk_notes": ["分數僅供研究參考。"],
        }
    )
    signals = pd.DataFrame(
        [{"date": "2024-01-01", "symbol": "2330", "signal": 1, "quantity": 10}]
    )
    result = BacktestEngine(initial_cash=10_000).run(prices, signals)
    alerts = [
        RiskAlert(
            code="max_drawdown_exceeded",
            severity="warning",
            message="回撤超過門檻。",
            symbol="2330",
        )
    ]
    return ReportData(
        symbol="2330",
        analysis_date="2024-01-04",
        price_data=prices,
        technical_indicators=indicators,
        fundamental_scores=fundamentals,
        backtest_result=result,
        risk_alerts=alerts,
        parameters={"strategy": "sample", "execution_price_col": "open"},
    )


def test_generate_excel_report_creates_required_sheets(tmp_path: Path) -> None:
    output_path = tmp_path / "report.xlsx"

    result_path = generate_excel_report(_report_data(), output_path)

    assert result_path.exists()
    workbook = load_workbook(result_path, read_only=True)
    assert workbook.sheetnames == list(SHEET_NAMES)
    summary = workbook["摘要"]
    values = [cell.value for row in summary.iter_rows() for cell in row if cell.value]
    assert "2330" in values
    assert any("歷史績效不代表未來報酬" in str(value) for value in values)
    risk_alert_values = [
        cell.value for row in workbook["風險提示"].iter_rows() for cell in row if cell.value
    ]
    assert SURVIVORSHIP_BIAS_WARNING in risk_alert_values
    assert TRAILING_STOP_DAILY_MODEL_NOTICE not in risk_alert_values
    assert workbook["交易紀錄"].max_row >= 2
    workbook.close()


def test_generate_html_report_contains_required_sections(tmp_path: Path) -> None:
    output_path = tmp_path / "report.html"

    result_path = generate_html_report(_report_data(), output_path)

    assert result_path.exists()
    html_text = result_path.read_text(encoding="utf-8")
    assert "歷史績效不代表未來報酬" in html_text
    assert "資金曲線" in html_text
    assert "回撤曲線" in html_text
    assert "交易紀錄" in html_text
    assert "指標摘要" in html_text
    assert "風險提示" in html_text
    assert SURVIVORSHIP_BIAS_WARNING in html_text
    assert TRAILING_STOP_DAILY_MODEL_NOTICE not in html_text
    assert "<svg" in html_text


def test_report_displays_insufficient_data_when_inputs_are_missing() -> None:
    report_data = ReportData(symbol="2330", analysis_date="2024-01-04")

    html_text = render_html_report(report_data)

    assert "資料不足" in html_text


def test_report_warns_when_point_in_time_universe_is_missing() -> None:
    report_data = _report_data()

    alerts = report_data.resolved_risk_alerts()

    assert SURVIVORSHIP_BIAS_WARNING in set(alerts["訊息"])


def test_report_marks_trailing_stop_daily_model_when_enabled() -> None:
    base = _report_data()
    report_data = ReportData(
        symbol=base.symbol,
        analysis_date=base.analysis_date,
        price_data=base.price_data,
        backtest_result=base.backtest_result,
        parameters={"strategy": "sample", "trailing_stop_pct": 0.10},
    )

    alerts = report_data.resolved_risk_alerts()

    assert TRAILING_STOP_DAILY_MODEL_NOTICE in set(alerts["訊息"])


def test_report_omits_survivorship_warning_when_universe_is_provided() -> None:
    base = _report_data()
    report_data = ReportData(
        symbol=base.symbol,
        analysis_date=base.analysis_date,
        price_data=base.price_data,
        backtest_result=base.backtest_result,
        point_in_time_universe="data/raw/historical_universe.csv",
    )

    alerts = report_data.resolved_risk_alerts()

    assert SURVIVORSHIP_BIAS_WARNING not in set(alerts["訊息"])

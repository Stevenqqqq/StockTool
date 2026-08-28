"""Excel report generation."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from stock_tool.reports.models import (
    DISCLAIMER,
    INSUFFICIENT_DATA,
    ReportData,
    ReportPackage,
    data_period,
    ensure_table,
    localize_display_frame,
    unpack_report_input,
)

SHEET_NAMES = (
    "摘要",
    "股價資料",
    "技術指標",
    "基本面評分",
    "回測結果",
    "交易紀錄",
    "資金曲線",
    "風險提示",
    "參數",
    "基準比較",
    "資料來源",
    "資料品質",
    "可重現性 Manifest",
    "公司行動稽核",
    "每日資產對帳",
)


def generate_excel_report(report_data: ReportData | ReportPackage, output_path: str | Path) -> Path:
    """Generate a multi-sheet Excel report and return the output path."""

    report_data = unpack_report_input(report_data)
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "摘要"
    _write_summary(summary, report_data)

    sheet_tables = {
        "股價資料": ensure_table(report_data.price_data, "股價資料"),
        "技術指標": ensure_table(
            report_data.technical_indicators,
            "技術指標",
        ),
        "基本面評分": ensure_table(
            report_data.fundamental_scores,
            "基本面評分",
        ),
        "回測結果": report_data.resolved_backtest_metrics(),
        "交易紀錄": report_data.resolved_trade_log(),
        "資金曲線": report_data.resolved_equity_curve(),
        "風險提示": report_data.resolved_risk_alerts(),
        "參數": report_data.resolved_parameters(),
        "基準比較": report_data.resolved_benchmark_comparison(),
        "資料來源": report_data.resolved_data_sources(),
        "資料品質": report_data.resolved_data_quality(),
        "可重現性 Manifest": report_data.resolved_manifest(),
        "公司行動稽核": report_data.resolved_corporate_action_audit(),
        "每日資產對帳": report_data.resolved_daily_reconciliation(),
    }

    for sheet_name in SHEET_NAMES[1:]:
        worksheet = workbook.create_sheet(title=sheet_name)
        _write_dataframe(worksheet, sheet_tables[sheet_name])

    workbook.save(path)
    return path


def _write_summary(worksheet: Worksheet, report_data: ReportData) -> None:
    worksheet["A1"] = "股票分析摘要"
    worksheet["A1"].font = Font(size=16, bold=True)
    worksheet.merge_cells("A1:D1")

    rows = [
        ("股票代號", report_data.symbol),
        ("分析日期", report_data.resolved_analysis_date()),
        ("資料期間", data_period(report_data.price_data)),
        ("技術面摘要", report_data.technical_summary or _technical_summary(report_data)),
        ("基本面摘要", report_data.fundamental_summary or _fundamental_summary(report_data)),
        ("風險摘要", report_data.risk_summary or _risk_summary(report_data)),
        ("回測績效摘要", report_data.backtest_summary or _backtest_summary(report_data)),
        ("重要聲明", DISCLAIMER),
    ]

    start_row = 3
    for index, (label, value) in enumerate(rows, start=start_row):
        worksheet.cell(index, 1, _excel_safe_value(label))
        worksheet.cell(index, 2, _excel_safe_value(value))
        worksheet.cell(index, 1).font = Font(bold=True)
        worksheet.cell(index, 2).alignment = Alignment(wrap_text=True, vertical="top")

    worksheet.column_dimensions["A"].width = 18
    worksheet.column_dimensions["B"].width = 96
    worksheet.freeze_panes = "A3"


def _write_dataframe(worksheet: Worksheet, frame: pd.DataFrame) -> None:
    table = localize_display_frame(frame)
    if table.empty:
        table = pd.DataFrame({"訊息": [INSUFFICIENT_DATA]})

    for column_index, column in enumerate(table.columns, start=1):
        cell = worksheet.cell(1, column_index, _excel_safe_value(str(column)))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
        cell.alignment = Alignment(horizontal="center")

    for row_index, (_, row) in enumerate(table.iterrows(), start=2):
        for column_index, column in enumerate(table.columns, start=1):
            value = row[column]
            cell = worksheet.cell(row_index, column_index, _excel_safe_value(value))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            _apply_number_format(cell, str(column))

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    _autosize_columns(worksheet)


def _autosize_columns(worksheet: Worksheet) -> None:
    for column_cells in worksheet.columns:
        max_length = 10
        column_letter = get_column_letter(column_cells[0].column)
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(value), 60))
        worksheet.column_dimensions[column_letter].width = max_length + 2


def _apply_number_format(cell: Any, column_name: str) -> None:
    lowered = column_name.lower()
    percent_tokens = (
        "return",
        "ratio",
        "rate",
        "yield",
        "drawdown",
        "exposure",
        "報酬",
        "比率",
        "比例",
        "率",
        "回撤",
        "曝險",
        "殖利率",
    )
    if any(token in lowered or token in column_name for token in percent_tokens):
        if isinstance(cell.value, (int, float)):
            cell.number_format = "0.00%"
    elif isinstance(cell.value, float):
        cell.number_format = "#,##0.00"


def _excel_safe_value(value: Any) -> Any:
    if isinstance(value, (list, tuple, set, dict)):
        value = str(value)
    if value is None:
        return INSUFFICIENT_DATA
    if isinstance(value, str):
        if value.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + value
        return value
    try:
        if bool(pd.isna(value)):
            return INSUFFICIENT_DATA
    except (TypeError, ValueError):
        pass
    if isinstance(value, float) and value == float("inf"):
        return "∞"
    return value


def _technical_summary(report_data: ReportData) -> str:
    indicators = report_data.technical_indicators
    if indicators is None or indicators.empty:
        return INSUFFICIENT_DATA
    latest = indicators.tail(1).iloc[0]
    parts: list[str] = []
    for column in ("close", "sma_20", "sma_60", "rsi_14", "macd_dif", "macd_dea"):
        if column in latest.index and not pd.isna(latest[column]):
            parts.append(
                f"{column}: {latest[column]:.2f}"
                if isinstance(latest[column], (int, float))
                else f"{column}: {latest[column]}"
            )
    return "；".join(parts) if parts else INSUFFICIENT_DATA


def _fundamental_summary(report_data: ReportData) -> str:
    scores = report_data.fundamental_scores
    if scores is None or scores.empty:
        return INSUFFICIENT_DATA
    row = scores.loc[scores["symbol"].astype(str) == str(report_data.symbol)]
    if row.empty:
        row = scores.head(1)
    latest = row.iloc[0]
    total = latest.get("total_score", INSUFFICIENT_DATA)
    strengths = latest.get("strengths", INSUFFICIENT_DATA)
    weaknesses = latest.get("weaknesses", INSUFFICIENT_DATA)
    return f"總分：{total}；優勢：{strengths}；弱項：{weaknesses}"


def _risk_summary(report_data: ReportData) -> str:
    alerts = report_data.resolved_risk_alerts()
    if alerts.empty:
        return INSUFFICIENT_DATA
    if "message" in alerts.columns and len(alerts) == 1:
        return str(alerts.iloc[0]["message"])
    if "訊息" in alerts.columns and len(alerts) == 1:
        return str(alerts.iloc[0]["訊息"])
    if "severity" in alerts.columns:
        counts = alerts["severity"].value_counts().to_dict()
        return "；".join(f"{severity}: {count}" for severity, count in counts.items())
    if "嚴重程度" in alerts.columns:
        counts = alerts["嚴重程度"].value_counts().to_dict()
        return "；".join(f"{severity}: {count}" for severity, count in counts.items())
    return f"{len(alerts)} 筆風險提示"


def _backtest_summary(report_data: ReportData) -> str:
    if report_data.backtest_result is None:
        return INSUFFICIENT_DATA
    metrics = report_data.backtest_result.metrics
    return (
        f"總報酬率：{metrics.total_return * 100:.2f}%；"
        f"年化報酬率：{metrics.cagr * 100:.2f}%；"
        f"最大回撤：{metrics.max_drawdown * 100:.2f}%；"
        f"交易次數：{metrics.number_of_trades}"
    )

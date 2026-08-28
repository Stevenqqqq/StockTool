"""Load stock price data from CSV and Excel files.

The loader only normalizes input columns into the project schema. Data quality
validation lives in ``stock_tool.data.cleaner`` so providers can be swapped
without duplicating validation logic.
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Iterable, Mapping

STANDARD_COLUMNS = (
    "date",
    "symbol",
    "open",
    "high",
    "low",
    "close",
    "volume",
    "adjusted_close",
)

REQUIRED_COLUMNS = (
    "date",
    "symbol",
    "open",
    "high",
    "low",
    "close",
    "volume",
)

COLUMN_ALIASES = {
    "date": "date",
    "datetime": "date",
    "trade_date": "date",
    "trading_date": "date",
    "日期": "date",
    "symbol": "symbol",
    "ticker": "symbol",
    "code": "symbol",
    "stock_id": "symbol",
    "股票代號": "symbol",
    "open": "open",
    "opening_price": "open",
    "開盤價": "open",
    "high": "high",
    "highest_price": "high",
    "最高價": "high",
    "low": "low",
    "lowest_price": "low",
    "最低價": "low",
    "close": "close",
    "closing_price": "close",
    "收盤價": "close",
    "volume": "volume",
    "成交量": "volume",
    "adjusted_close": "adjusted_close",
    "adj_close": "adjusted_close",
    "adjclose": "adjusted_close",
    "調整後收盤價": "adjusted_close",
}


class MissingColumnError(ValueError):
    """Raised when an input file does not contain required price columns."""

    def __init__(self, missing_columns: Iterable[str]) -> None:
        self.missing_columns = tuple(missing_columns)
        joined = ", ".join(self.missing_columns)
        super().__init__(f"缺少必要欄位：{joined}")


def load_csv(
    file_path: str | Path,
    *,
    encoding: str = "utf-8-sig",
    column_mapping: Mapping[str, str] | None = None,
) -> list[dict[str, Any]]:
    """Load a CSV file and normalize it to the standard price schema.

    Args:
        file_path: CSV file path.
        encoding: File encoding. ``utf-8-sig`` handles common BOM CSV files.
        column_mapping: Optional mapping from source column name to standard
            column name. It is applied before built-in aliases.

    Returns:
        A list of dictionaries with all standard columns present.

    Raises:
        MissingColumnError: If required columns are missing.
    """

    path = Path(file_path)
    with path.open("r", encoding=encoding, newline="") as csv_file:
        reader = csv.DictReader(csv_file)
        if reader.fieldnames is None:
            raise MissingColumnError(REQUIRED_COLUMNS)

        normalized_headers = _normalize_headers(reader.fieldnames, column_mapping)
        _ensure_required_columns(normalized_headers.values())

        rows: list[dict[str, Any]] = []
        for raw_row in reader:
            rows.append(_standardize_row(raw_row, normalized_headers))

    return rows


def load_excel(
    file_path: str | Path,
    *,
    sheet_name: str | None = None,
    column_mapping: Mapping[str, str] | None = None,
) -> list[dict[str, Any]]:
    """Load an Excel worksheet and normalize it to the standard price schema.

    Args:
        file_path: Excel file path.
        sheet_name: Worksheet name. If omitted, the active sheet is used.
        column_mapping: Optional mapping from source column name to standard
            column name. It is applied before built-in aliases.

    Returns:
        A list of dictionaries with all standard columns present.

    Raises:
        ImportError: If ``openpyxl`` is not installed.
        MissingColumnError: If required columns are missing.
    """

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("匯入 Excel 需要 openpyxl。請執行：pip install openpyxl") from exc

    workbook = load_workbook(filename=Path(file_path), read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active

    rows_iter = worksheet.iter_rows(values_only=True)
    try:
        header_values = next(rows_iter)
    except StopIteration as exc:
        raise MissingColumnError(REQUIRED_COLUMNS) from exc

    headers = ["" if value is None else str(value) for value in header_values]
    normalized_headers = _normalize_headers(headers, column_mapping)
    _ensure_required_columns(normalized_headers.values())

    rows: list[dict[str, Any]] = []
    for values in rows_iter:
        raw_row = {
            header: value
            for header, value in zip(headers, values)
            if header is not None and str(header).strip()
        }
        if all(value is None or str(value).strip() == "" for value in raw_row.values()):
            continue
        rows.append(_standardize_row(raw_row, normalized_headers))

    workbook.close()
    return rows


def _normalize_headers(
    headers: Iterable[str],
    column_mapping: Mapping[str, str] | None,
) -> dict[str, str]:
    normalized: dict[str, str] = {}
    explicit_mapping = {
        _normalize_name(source): target for source, target in (column_mapping or {}).items()
    }

    for header in headers:
        raw_header = str(header)
        normalized_name = _normalize_name(raw_header)
        mapped = explicit_mapping.get(normalized_name)
        if mapped is None:
            mapped = COLUMN_ALIASES.get(normalized_name, normalized_name)

        if mapped in STANDARD_COLUMNS:
            normalized[raw_header] = mapped

    return normalized


def _standardize_row(
    raw_row: Mapping[str, Any],
    normalized_headers: Mapping[str, str],
) -> dict[str, Any]:
    row = {column: None for column in STANDARD_COLUMNS}
    for source_column, standard_column in normalized_headers.items():
        row[standard_column] = raw_row.get(source_column)

    return row


def _ensure_required_columns(columns: Iterable[str]) -> None:
    present = set(columns)
    missing = [column for column in REQUIRED_COLUMNS if column not in present]
    if missing:
        raise MissingColumnError(missing)


def _normalize_name(name: str) -> str:
    return name.strip().lower().replace(" ", "_").replace("-", "_")
